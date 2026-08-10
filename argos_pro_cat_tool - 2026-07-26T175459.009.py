import json
import os
import threading
import xml.etree.ElementTree as ET
from datetime import datetime
from tkinter import filedialog, messagebox
import re  # Додано для регулярних виразів

import customtkinter as ctk
import pandas as pd
import spacy
from docx import Document

# Інтеграція Argos Translate
try:
    import argostranslate.package
    import argostranslate.translate

    ARGOS_AVAILABLE = True
except ImportError:
    ARGOS_AVAILABLE = False

# Завантаження NLP моделей spaCy для виділення термінів
nlp_models = {}
try:
    nlp_models["en"] = spacy.load("en_core_web_sm")
except Exception:
    nlp_models["en"] = None

try:
    nlp_models["uk"] = spacy.load("uk_core_news_sm")
except Exception:
    nlp_models["uk"] = None

# Глобальні дані програми
glossary = {}  # { "source_term": "target_term" }
translation_history = []  # Накопичувальна пам'ять перекладів (Translation Memory)
loaded_docx_path = None  # Шлях до завантаженого DOCX документу
loaded_docx_obj = None  # Об'єкт Document для збереження форматування

AUTO_SAVE_FILE = "auto_translation_history.json"
PERSISTENT_TMX_FILE = "translation_memory.tmx"

def load_auto_save():
    global translation_history
    translation_history = []
    if os.path.exists(AUTO_SAVE_FILE):
        try:
            with open(AUTO_SAVE_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list): translation_history.extend(data)
        except Exception as e: print(f"Помилка завантаження автозбереження: {e}")
    if os.path.exists(PERSISTENT_TMX_FILE): import_tmx_file(PERSISTENT_TMX_FILE, silent=True)

def auto_save_translation(src_text, tgt_text, src_lang, tgt_lang):
    global translation_history
    if not src_text.strip() or not tgt_text.strip(): return
    entry = {"timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "source_lang": src_lang, "target_lang": tgt_lang, "source_text": src_text, "target_text": tgt_text}
    if entry not in translation_history: translation_history.append(entry)
    try:
        with open(AUTO_SAVE_FILE, "w", encoding="utf-8") as f: json.dump(translation_history, f, ensure_ascii=False, indent=2)
    except Exception as e: print(f"Помилка автозбереження: {e}")

def load_file():
    global loaded_docx_path, loaded_docx_obj
    path = filedialog.askopenfilename(filetypes=[("Документи (*.docx, *.txt)", "*.docx *.txt")])
    if not path: return
    try:
        if path.endswith(".docx"):
            loaded_docx_path = path
            loaded_docx_obj = Document(path)
            full_text = "\n".join([p.text for p in loaded_docx_obj.paragraphs])
            btn_translate_doc.configure(state="normal")
        else:
            loaded_docx_path = None
            loaded_docx_obj = None
            btn_translate_doc.configure(state="disabled")
            with open(path, "r", encoding="utf-8") as f: full_text = f.read()
        input_area.delete("1.0", "end")
        input_area.insert("1.0", full_text)
        status_label.configure(text=f"Файл: {os.path.basename(path)}")
    except Exception as e: messagebox.showerror("Помилка", f"Не вдалося відкрити файл: {e}")

def setup_argos_language_pair(from_code, to_code):
    if not ARGOS_AVAILABLE: return False
    try:
        installed_langs = argostranslate.translate.get_installed_languages()
        src_lang = next((l for l in installed_langs if l.code == from_code), None)
        tgt_lang = next((l for l in installed_langs if l.code == to_code), None)
        if src_lang and tgt_lang and src_lang.get_translation(tgt_lang): return True
        status_label.configure(text=f"Завантаження моделі {from_code}➔{to_code}...")
        root.update_idletasks()
        argostranslate.package.update_package_index()
        available_packages = argostranslate.package.get_available_packages()
        package_to_install = next((p for p in available_packages if p.from_code == from_code and p.to_code == to_code), None)
        if package_to_install:
            download_path = package_to_install.download()
            argostranslate.package.install_from_path(download_path)
            return True
        return False
    except Exception as e:
        print(f"Помилка Argos Translate: {e}")
        return False

def apply_glossary(text, glossary_dict):
    """
    Професійна заміна термінів за допомогою регулярних виразів.
    Замінює лише цілі слова (word boundaries) та ігнорує регістр при пошуку.
    """
    if not glossary_dict:
        return text
        
    for src_term, tgt_term in glossary_dict.items():
        if not src_term or not str(tgt_term).strip() or str(tgt_term).lower() == "nan":
            continue
        # \b - межа слова, re.IGNORECASE - ігнорування регістру
        # re.escape - захищає від спецсимволів у термінах
        pattern = re.compile(r'\b' + re.escape(src_term) + r'\b', re.IGNORECASE)
        text = pattern.sub(str(tgt_term), text)
    return text

def translate_text_segment(text, src_lang, tgt_lang):
    if not text or not text.strip(): return text
    
    # 1. Пошук у Пам'яті перекладів (TM) - 100% Match
    for entry in translation_history:
        if entry.get("source_lang") == src_lang and entry.get("target_lang") == tgt_lang and entry.get("source_text") == text.strip():
            return entry.get("target_text")
            
    # 2. Машинний переклад (MT)
    translated_text = text
    if ARGOS_AVAILABLE:
        try:
            setup_argos_language_pair(src_lang, tgt_lang)
            translated_text = argostranslate.translate.translate(text, src_lang, tgt_lang)
        except Exception as e: 
            print(f"Помилка перекладу: {e}")
            
    # 3. Накладання глосарію на результат перекладу
    final_text = apply_glossary(translated_text, glossary)
    
    return final_text

def extract_terms():
    global glossary
    src_lang, tgt_lang = [code.lower() for code in lang_pair_var.get().split(" ➔ ")]
    text = input_area.get("1.0", "end-1c").strip()
    if not text:
        messagebox.showwarning("Увага", "Спочатку завантажте текст!")
        return
    nlp = nlp_models.get(src_lang)
    if not nlp:
        messagebox.showerror("Помилка", f"NLP модель '{src_lang}' не завантажена!")
        return
    status_label.configure(text="Виділення термінів...")
    root.update_idletasks()
    if ARGOS_AVAILABLE: setup_argos_language_pair(src_lang, tgt_lang)
    doc = nlp(text)
    stopwords = {"a", "an", "the", "her", "his", "my", "their", "our", "its"}
    raw_terms = [chunk.text.lower() for chunk in doc.noun_chunks if 1 <= len(chunk.text.split()) <= 4]
    cleaned_terms = set()
    for term in raw_terms:
        words = term.split()
        if words and words[0] in stopwords: words = words[1:]
        if words: cleaned_terms.add(" ".join(words))
    for term in cleaned_terms:
        if term not in glossary or not glossary[term]:
            translated_term = term
            if ARGOS_AVAILABLE:
                try: translated_term = argostranslate.translate.translate(term, src_lang, tgt_lang)
                except Exception as e: print(f"Помилка перекладу терміна: {e}")
            glossary[term] = translated_term
    messagebox.showinfo("Успіх", "Терміни виділено та перекладено.")

def import_glossary():
    global glossary
    path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
    if not path: return
    try:
        df = pd.read_excel(path)
        for _, row in df.iterrows():
            if pd.notna(row[0]): glossary[str(row[0]).strip().lower()] = str(row[1]) if len(row) > 1 else ""
        messagebox.showinfo("Успіх", "Словник імпортовано.")
    except Exception as e: messagebox.showerror("Помилка", f"Не вдалося прочитати: {e}")

def export_glossary():
    path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
    if not path: return
    pd.DataFrame(list(glossary.items()), columns=["Source", "Target"]).to_excel(path, index=False)
    messagebox.showinfo("Успіх", "Словник збережено.")

def import_tmx_file(path=None, silent=False):
    global translation_history
    if not path: path = filedialog.askopenfilename(filetypes=[("TMX", "*.tmx")])
    if not path or not os.path.exists(path): return
    try:
        tree = ET.parse(path)
        for tu in tree.getroot().iter("tu"):
            tuvs = list(tu.iter("tuv"))
            if len(tuvs) >= 2:
                txt1 = tuvs[0].find("seg").text if tuvs[0].find("seg") is not None else ""
                txt2 = tuvs[1].find("seg").text if tuvs[1].find("seg") is not None else ""
                if txt1 and txt2:
                    entry = {"source_lang": "en", "target_lang": "uk", "source_text": txt1, "target_text": txt2}
                    if entry not in translation_history: translation_history.append(entry)
        if not silent: messagebox.showinfo("Успіх", "TMX імпортовано.")
    except Exception as e: messagebox.showerror("Помилка", f"TMX файл помилка: {e}")

def export_tmx():
    path = filedialog.asksaveasfilename(defaultextension=".tmx", filetypes=[("TMX", "*.tmx")])
    if not path: return
    root_el = ET.Element("tmx", version="1.4")
    body = ET.SubElement(ET.SubElement(root_el, "header"), "body")
    for item in translation_history:
        tu = ET.SubElement(body, "tu")
        ET.SubElement(ET.SubElement(tu, "tuv", xml_lang=item["source_lang"]), "seg").text = item["source_text"]
        ET.SubElement(ET.SubElement(tu, "tuv", xml_lang=item["target_lang"]), "seg").text = item["target_text"]
    tree = ET.ElementTree(root_el)
    ET.indent(tree, space="  ")
    tree.write(path, encoding="utf-8", xml_declaration=True)
    messagebox.showinfo("Успіх", "TMX збережено.")

def save_docx():
    path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word", "*.docx")])
    if not path: return
    doc = Document()
    for paragraph in output_area.get("1.0", "end-1c").split("\n"):
        if paragraph.strip(): doc.add_paragraph(paragraph)
    doc.save(path)
    messagebox.showinfo("Успіх", "Документ збережено.")

def run_document_translation_thread():
    global loaded_docx_obj
    src_lang, tgt_lang = [code.lower() for code in lang_pair_var.get().split(" ➔ ")]
    if not loaded_docx_obj: return
    save_path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word", "*.docx")])
    if not save_path: return
    
    def process_docx(doc_obj, src, tgt):
        # Покращена логіка: перекладаємо цілий абзац, щоб не розривати контекст
        # Зберігаємо стиль через перший Run
        for p in doc_obj.paragraphs:
            if not p.text.strip(): continue
            
            translated = translate_text_segment(p.text, src, tgt)
            if p.runs:
                p.runs[0].text = translated
                for run in p.runs[1:]: 
                    run.text = ""
            else:
                p.add_run(translated)
                
        # Переклад вмісту таблиць
        for table in doc_obj.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if not p.text.strip(): continue
                        translated = translate_text_segment(p.text, src, tgt)
                        if p.runs:
                            p.runs[0].text = translated
                            for run in p.runs[1:]: 
                                run.text = ""
                        else:
                            p.add_run(translated)
        return doc_obj

    try:
        translated_doc = process_docx(loaded_docx_obj, src_lang, tgt_lang)
        translated_doc.save(save_path)
        messagebox.showinfo("Успіх", "Документ збережено!")
    except Exception as e: messagebox.showerror("Помилка", f"{e}")

def run_translation_thread():
    src_lang, tgt_lang = [code.lower() for code in lang_pair_var.get().split(" ➔ ")]
    src_text = input_area.get("1.0", "end-1c").strip()
    if not src_text: return
    
    # Виправлено помилку з дублюванням функції перекладу
    translated_text = translate_text_segment(src_text, src_lang, tgt_lang)
    
    output_area.delete("1.0", "end")
    output_area.insert("1.0", translated_text)
    auto_save_translation(src_text, translated_text, src_lang, tgt_lang)

load_auto_save()
ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")
root = ctk.CTk()
root.title("Argos Pro CAT Tool")
root.geometry("1050x950")

# Header Frame
top_frame = ctk.CTkFrame(root)
top_frame.pack(pady=10, padx=20, fill="x")
lang_pair_var = ctk.StringVar(value="EN ➔ UK")
ctk.CTkOptionMenu(top_frame, values=["EN ➔ UK", "UK ➔ EN"], variable=lang_pair_var).pack(side="left", padx=10, pady=10)
btn_load_file = ctk.CTkButton(top_frame, text="📁 Завантажити DOCX/TXT", command=load_file)
btn_load_file.pack(side="left", padx=10)
btn_translate_doc = ctk.CTkButton(top_frame, text="📄 Перекласти DOCX (Формат)", command=lambda: threading.Thread(target=run_document_translation_thread, daemon=True).start(), fg_color="#2b8a3e")
btn_translate_doc.pack(side="left", padx=10)

# Input Area
ctk.CTkLabel(root, text="Source Text:", font=("Arial", 12, "bold")).pack(anchor="w", padx=20)
input_area = ctk.CTkTextbox(root, height=180)
input_area.pack(pady=5, padx=20, fill="x")

# ACTION PANEL (Translation)
action_panel = ctk.CTkFrame(root, fg_color="transparent")
action_panel.pack(pady=10, padx=20, fill="x")
ctk.CTkButton(action_panel, text="🔍 Виділити терміни (NLP)", command=extract_terms, fg_color="#15aabf").pack(side="left", padx=5)
ctk.CTkButton(action_panel, text="⚡ Перекласти текст", command=lambda: threading.Thread(target=run_translation_thread, daemon=True).start(), fg_color="green").pack(side="left", padx=5)

# STORAGE/EXPORT PANEL (Save Buttons Visualized)
storage_panel = ctk.CTkFrame(root, border_width=2)
storage_panel.pack(pady=10, padx=20, fill="x")
ctk.CTkLabel(storage_panel, text="💾 Операції Збереження та Експорту", font=("Arial", 12, "bold")).pack(pady=5)

storage_grid = ctk.CTkFrame(storage_panel, fg_color="transparent")
storage_grid.pack(pady=5)

ctk.CTkButton(storage_grid, text="💾 Зберегти (DOCX)", command=save_docx, fg_color="#d9480f").grid(row=0, column=0, padx=5, pady=5)
ctk.CTkButton(storage_grid, text="📖 Імпорт словника (XLSX)", command=import_glossary, fg_color="#4c6ef5").grid(row=0, column=1, padx=5, pady=5)
ctk.CTkButton(storage_grid, text="📊 Експорт словника (XLSX)", command=export_glossary, fg_color="#4c6ef5").grid(row=0, column=2, padx=5, pady=5)
ctk.CTkButton(storage_grid, text="📥 Імпорт пам'яті (TMX)", command=import_tmx_file, fg_color="#ae3ec9").grid(row=1, column=0, padx=5, pady=5)
ctk.CTkButton(storage_grid, text="📤 Експорт пам'яті (TMX)", command=export_tmx, fg_color="#ae3ec9").grid(row=1, column=1, padx=5, pady=5)

# Output Area
ctk.CTkLabel(root, text="Target Text:", font=("Arial", 12, "bold")).pack(anchor="w", padx=20)
output_area = ctk.CTkTextbox(root, height=180)
output_area.pack(pady=5, padx=20, fill="x")

status_label = ctk.CTkLabel(root, text="Ready")
status_label.pack(pady=5)

root.mainloop()